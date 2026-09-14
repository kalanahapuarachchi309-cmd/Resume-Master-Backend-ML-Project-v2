"""DOCX Document Parser Service (Mahen & Team)."""
import io
import zipfile
import xml.etree.ElementTree as ET
import logging
from typing import Optional

logger = logging.getLogger(__name__)


class DocxParser:
    """Extracts text content from Microsoft Word (.docx) resume documents."""

    @staticmethod
    def validate_file(file_bytes: bytes) -> bool:
        """Verify that the byte stream is a valid DOCX zip archive header (PK\\x03\\x04)."""
        return file_bytes.startswith(b"PK\x03\x04")

    @classmethod
    def extract_text(cls, file_bytes: bytes) -> str:
        """Extract paragraph text and table contents from a DOCX byte buffer."""
        if not file_bytes:
            return ""

        # 1. Try python-docx if installed
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        full_text.append(row_text)
            if full_text:
                return "\n".join(full_text)
        except Exception as e:
            logger.debug(f"python-docx extraction skipped: {e}")

        # 2. Pure Python XML Fallback from zip archive
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as docx_zip:
                xml_content = docx_zip.read("word/document.xml")
                tree = ET.fromstring(xml_content)
                # Word XML namespace for text elements
                namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                text_pieces = []
                for node in tree.iterfind(".//w:t", namespaces):
                    if node.text:
                        text_pieces.append(node.text)
                return " ".join(text_pieces).strip()
        except Exception as e:
            logger.error(f"XML docx extraction error: {e}")

        return ""
